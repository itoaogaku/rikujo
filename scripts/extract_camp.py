"""合宿の走行距離・消化率Excelから選手別の整形CSVを作る
usage: python3 scripts/extract_camp.py <走行距離.xlsx> <消化率.xlsx>"""
import csv, re, sys
import openpyxl

def norm(n):
    return re.sub(r'\s', '', str(n)).replace('熙', '煕').replace('昴', '昂')

def num(v):
    return v if isinstance(v, (int, float)) else ''

def table(ws, ncols):
    """'名前','1次合宿',... の見出しから下へ、名前と続くncols列を読む"""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            if c == '名前' and j + 1 < len(r) and r[j + 1] == '1次合宿':
                out = []
                for rr in rows[i + 1:]:
                    n = rr[j]
                    if not n or '平均' in str(n) or 'チーム' in str(n):
                        break
                    out.append([norm(n)] + [num(x) for x in rr[j + 1:j + 1 + ncols]])
                return out
    raise ValueError(ws.title)

dist_wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
with open('data/camp_distance.csv', 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['year', 'name', 'camp1_km', 'camp2_km'])
    for year, sheet in [('2024', '１次２次まとめ'), ('2025', '2025１次２次まとめ')]:
        for r in table(dist_wb[sheet], 2):
            w.writerow([year] + r)

comp_wb = openpyxl.load_workbook(sys.argv[2], data_only=True)
with open('data/camp_completion.csv', 'w', newline='') as f:
    w = csv.writer(f)
    w.writerow(['year', 'name', 'camp1', 'camp2', 'camp3', 'summer_avg'])
    for year, sheet in [('2023', '合宿消化率まとめ'), ('2025', '2025合宿消化率まとめ')]:
        for r in table(comp_wb[sheet], 4):
            w.writerow([year] + [r[0]] + [round(x, 3) if x != '' else '' for x in r[1:]])
