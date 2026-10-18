"""返ってきたスカウト一覧Excelから、オレンジのマーカー（注目選手）とメモ列を data/scout_marks_2028.csv に取り込む
（どのシートでマークしても可。選手は各シートのNo列から選手一覧の氏名・高校を引いて記録する）
usage: python3 scripts/import_marks.py 返ってきた.xlsx data/scout_marks_2028.csv"""
import csv, sys
from openpyxl import load_workbook

ORANGE = 'FFFFC000'
wb = load_workbook(sys.argv[1])
ws = wb['選手一覧']
head = [c.value for c in ws[4]]
ci = {h: i for i, h in enumerate(head)}
players = {row[0]: row for row in ws.iter_rows(min_row=5, values_only=True) if row[0]}
marks = {}  # No -> [focus, memo]
for row in ws.iter_rows(min_row=5):
    no, memo = row[0].value, row[ci['メモ']].value
    if row[1].fill.fgColor.rgb == ORANGE or memo:
        marks[no] = [row[1].fill.fgColor.rgb == ORANGE, memo or '']
for s in wb.worksheets[1:]:
    hs = [c.value for c in s[4]]
    if 'No' not in hs:
        continue
    k = hs.index('No')
    for row in s.iter_rows(min_row=5):
        if row[k].value and any(c.fill.fgColor.rgb == ORANGE for c in row):
            marks.setdefault(row[k].value, [True, ''])[0] = True
w = csv.writer(open(sys.argv[2], 'w', encoding='utf-8', newline=''))
w.writerow(['name', 'school', 'focus', 'memo'])
for no, (focus, memo) in sorted(marks.items()):
    p = players[no]
    w.writerow([p[1], p[2], int(focus), memo])
print(len(marks), 'players,', sum(f for f, _ in marks.values()), 'focus')
