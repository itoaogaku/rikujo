"""陸マガの高校ランキング（data/ranking_2028.xlsx：5000m・3000mSC・1500mが横並びの1シート）から
選手単位にまとめたスカウト用Excelを作る。5000mを主に、並び順・集計は5000mの順位を基準にする。
駅伝の区間成績（RACESのCSV）も、ランキングに載っている選手（＝2028年入学世代）だけ照合して追記する
usage: python3 scripts/build_scout_ranking.py data/ranking_2028.xlsx output/2028入学世代_スカウト一覧.xlsx"""
import csv, re, sys, unicodedata
from bisect import bisect_left
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as col_idx, get_column_letter

FONT = 'HGPｺﾞｼｯｸM'
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FONT = Font(name=FONT, bold=True, color='000000', size=10)
HEAD_FILL = PatternFill('solid', fgColor='A9D08E')
INPUT_FILL = PatternFill('solid', fgColor='FFFF99')
BASE = Font(name=FONT, size=10, color='000000')
BOLD = Font(name=FONT, size=10, color='000000', bold=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=False)
TFMT = '[m]:ss.00'
TOP10 = PatternFill('solid', fgColor='63BE7B')
TOP50 = PatternFill('solid', fgColor='A9D08E')
TOP100 = PatternFill('solid', fgColor='E2EFDA')

EVENTS = ['5000m', '3000mSC', '1500m']
# 同一人物で表記が違うもの（同じ高校・同じ種目に二重掲載されているもの含む）
NAME_ALIAS = {'ポールンジョロゲ': 'ポールジョロゲ', 'PAREIYOANTONYITA': 'パレイヨアントニーイタ',
              # 外字が化けたもの・欠けたもの（いずれも同じ高校）
              '??光寺奏': '座光寺奏', '??村悠希': '吉村悠希', '譲次郎': '高柳譲次郎', '柳譲次郎': '高柳譲次郎'}
# 陸マガ5000mランキングの新しい版（サイトからコピーしたテキスト。高校名なし）。速い記録だけ上書きする
R5000_UPDATE = 'data/ranking5000_20261018.txt'
# 同名で高校が違っても同一人物とみなす下限（校名の文字の重なり率）
SAME_SCHOOL = 0.6
# 駅伝：(シート名, 区間成績CSV, 大会名)。CSVは scripts/extract_nihonkai.py・extract_okukuma.py で作る
RACES = [('日本海駅伝', 'data/nihonkai_2026.csv', '第46回 日本海駅伝（2026/10/4）'),
         ('奥球磨駅伝', 'data/okukuma_2026.csv', '第5回 奥球磨駅伝 高校の部（2026/10/4）')]
# 旧字体など（陸マガのランキングは使っていない字）を揃える
KANJI = str.maketrans('﨑髙濵澤邊邉', '崎高浜沢辺辺')
# 注目選手（オレンジ）とメモ。返ってきたExcelから scripts/import_marks.py で取り込む
MARKS = 'data/scout_marks_2028.csv'
FOCUS_FILL = PatternFill('solid', fgColor='FFC000')
MAIN_FILL = PatternFill('solid', fgColor='548235')
MAIN_HEAD = Font(name=FONT, bold=True, color='FFFFFF', size=10)


def key(name):
    k = re.sub(r'\s', '', unicodedata.normalize('NFKC', name)).translate(KANJI)
    return NAME_ALIAS.get(k, k)


def school_core(s):
    s = unicodedata.normalize('NFKC', s)
    return re.sub(r'(高校|高|陸協)$', '', s.split('・')[0])


def similar(a, b):
    a, b = school_core(a), school_core(b)
    if not a or not b:
        return False
    return len(set(a) & set(b)) / min(len(set(a)), len(set(b))) >= SAME_SCHOOL


def tval(s):
    """'mm.ss.xx'（陸マガ表記）を日数（Excelの時刻値）に"""
    m, sec, frac = s.split('.')
    return (int(m) * 60 + int(sec) + int(frac) / 10 ** len(frac)) / 86400


def put(ws, r, c, v, fmt=None, font=BASE, align=CENTER):
    cell = ws.cell(r, c, v)
    cell.font, cell.border, cell.alignment = font, BORDER, align
    if fmt:
        cell.number_format = fmt
    return cell


def header(ws, r, labels, widths):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(r, j, lab)
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, CENTER, BORDER
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1]
    ws.row_dimensions[r].height = 14 * max(l.count('\n') + 1 for l in labels) + 6


def team_schools(team):
    """駅伝のチーム名（例：佐久長聖A、鹿児島実業高校B、小林・延岡工業高校合同）を校名の候補に"""
    team = re.sub(r'(合同|[A-ZＡ-Ｚ])$', '', team)
    return [re.sub(r'(高校|高)$', '', t) + '高' for t in team.split('・')]


def time_val(s):
    """'h:mm:ss' を日数に"""
    h, m, sec = map(int, s.split(':'))
    return (h * 3600 + m * 60 + sec) / 86400


def title(ws, text, sub):
    ws['A1'] = text
    ws['A1'].font = Font(name=FONT, bold=True, size=14, color='000000')
    ws['A2'] = sub
    ws['A2'].font = Font(name=FONT, size=9, color='000000')


def rank_colors(ws, ref, steps=((10, TOP10), (50, TOP50), (100, TOP100))):
    """順位セルの色分け。空欄（記録なし）は0扱いで色が付くので、数値のセルだけに限る"""
    for rg in ref.split():
        top = rg.split(':')[0]
        for val, fill in steps:
            ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({top}),{top}<={val})'], fill=fill,
                                                          stopIfTrue=True))


# ---------------- 読み込み ----------------
src = load_workbook(sys.argv[1]).active
head = [c.value for c in src[1]]
entries = []  # (ev, rank, name, school, time_str)
for row in src.iter_rows(min_row=2, values_only=True):
    for ev in EVENTS:
        i = head.index(f'{ev}_順位')
        rk, name, school, t = row[i:i + 4]
        if name:
            entries.append((ev, int(rk), str(name), str(school), str(t)))

# ---------------- 選手単位に統合 ----------------
athletes = []  # dict(name, schools, ev->(rank,time), notes)
by_key = defaultdict(list)
for ev, rk, name, school, t in entries:
    k = key(name)
    a = next((a for a in by_key[k] if any(similar(school, s) for s in a['schools'])), None)
    if a is None:
        a = {'name': name, 'schools': [], 'ev': {}, 'src': {}, 'notes': []}
        by_key[k].append(a)
        athletes.append(a)
    if school not in a['schools']:
        a['schools'].append(school)
    raw = lambda n: re.sub(r'\s', '', unicodedata.normalize('NFKC', n)).translate(KANJI)
    if raw(name) == k and raw(a['name']) != k:
        a['name'] = name  # 化けた・欠けた・ローマ字の表記より、NAME_ALIASの正しい表記を優先
    if ev in a['ev'] and rk > a['ev'][ev][0]:
        continue  # 同一人物の二重掲載は上位の記録を採用
    a['ev'][ev] = (rk, t)
    a['src'][ev] = (name, school)  # 種目別シートに出す元データの表記
# ---------------- 5000mの新しいランキングで、速くなった記録を更新 ----------------
txt = open(R5000_UPDATE, encoding='utf-8').read()
newest = {}  # key -> (name, time, date)  同じ名前が複数あれば最も速い記録
for rk, t, name, _, date in re.findall(r'(?m)^(\d+)\t(\d+\.\d\d\.\d\d)\t(.*?)\s*(\d)\t(\d\d/\d\d/\d\d)', txt, re.S):
    name = name.replace('\n', '')
    if key(name) not in newest or tval(t) < tval(newest[key(name)][1]):
        newest[key(name)] = (name, t, date)
for k, (name, t, date) in newest.items():
    group = by_key.get(k, [])
    with5 = [a for a in group if '5000m' in a['ev']]
    if len(group) > 1 and len(with5) != 1:
        continue  # 同名が複数いてどの選手か決められない
    if not group:  # 旧ランキングにいない選手（高校名は不明）
        a = {'name': name, 'schools': [''], 'ev': {}, 'src': {}, 'notes': ['高校名不明（5000m新ランキングのみ）']}
        by_key[k].append(a)
        athletes.append(a)
        group = [a]
    a = (with5 or group)[0]
    old = a['ev'].get('5000m')
    if old is None or tval(t) < tval(old[1]):
        a['ev']['5000m'] = (None, t)
        a['src']['5000m'] = (name, a['schools'][0])
        a['notes'].append(f'5000m更新：{t}（20{date}）' + (f'←{old[1]}' if old else ''))
# 5000mの順位は記録順に振り直す（同タイムは同順位）
times = sorted(tval(a['ev']['5000m'][1]) for a in athletes if '5000m' in a['ev'])
for a in athletes:
    if '5000m' in a['ev']:
        a['ev']['5000m'] = (bisect_left(times, tval(a['ev']['5000m'][1]) - 1e-9) + 1, a['ev']['5000m'][1])

for k, group in by_key.items():
    if len(group) > 1:
        for a in group:
            others = '・'.join(s for b in group if b is not a for s in b['schools'])
            a['notes'].append(f'同名の別選手あり（{others}）')
for a in athletes:
    # 表示用の校名：「高」で終わり、中黒を含まない最長の表記
    cands = sorted(a['schools'], key=lambda s: (s.endswith('高'), '・' not in s, len(s)), reverse=True)
    a['school'] = cands[0]
    if len(a['schools']) > 1:
        a['notes'].append('校名表記：' + '／'.join(a['schools']))
    hand = [ev for ev, (_, t) in a['ev'].items() if len(t.split('.')[-1]) == 1]
    if hand:
        a['notes'].append('手動計時：' + '・'.join(hand))
    a['best'] = min(r for r, _ in a['ev'].values())
    a['r5000'] = a['ev'].get('5000m', (None,))[0]
# 並び順：5000mの順位。5000mの記録がない選手はその後ろに、他種目の最高順位順
athletes.sort(key=lambda a: (a['r5000'] is None, a['r5000'] or a['best'], -len(a['ev'])))
NO = {id(a): i for i, a in enumerate(athletes, 1)}
for m in csv.DictReader(open(MARKS, encoding='utf-8')):
    a = next(a for a in by_key[key(m['name'])] if any(similar(m['school'], s) for s in a['schools']))
    a['focus'], a['memo'] = m['focus'] == '1', m['memo'] or None
FOCUS = [a for a in athletes if a.get('focus')]


def mark_row(ws, r, ncol, a):
    """注目選手の行をオレンジにする"""
    if a.get('focus'):
        for c in range(1, ncol + 1):
            ws.cell(r, c).fill = FOCUS_FILL

# ---------------- 駅伝：名前が同じで、チーム名と校名が似ている選手だけ照合 ----------------
races = []
for sheet, path, label in RACES:
    runs = []
    for x in csv.DictReader(open(path, encoding='utf-8')):
        a = next((a for a in by_key.get(key(x['name']), [])
                  if any(similar(t, s) for t in team_schools(x['team']) for s in a['schools'])), None)
        if a:
            x['no'] = NO[id(a)]
            a.setdefault('races', {})[sheet] = f"{x['section']}区 {x['rank']}位（{x['time'].lstrip('0:')}）"
            runs.append(x)
    runs.sort(key=lambda x: (int(x['section']), int(x['rank'])))
    races.append((sheet, label, runs))


wb = Workbook()

# ---------------- 選手一覧 ----------------
ws = wb.active
ws.title = '選手一覧'
title(ws, '2028年入学世代（現高校2年）中長距離 選手一覧（5000m順）',
      '陸マガの5000m・3000mSC・1500mランキングを選手ごとに1行にまとめたもの。5000mは10/18時点のランキングで速い記録に更新し、'
      '順位は記録順に振り直し。並び順＝5000mの順位（5000mの記録がない選手は'
      'その後ろに他種目の最高順位順）。順位の色：濃い緑＝10位以内、緑＝50位以内、薄緑＝100位以内。'
      '見出しのフィルターで高校名・駅伝出場などを絞り込めます。オレンジ＝注目選手（「注目選手」シートに一覧）。'
      'メモ列（黄色）は自由に記入してください。')
H = ['No', '氏名', '高校', '5000m\n記録', '5000m\n順位', '3000mSC\n記録', '3000mSC\n順位', '1500m\n記録', '1500m\n順位']
H += [f'{sheet}\n（区・区間順位）' for sheet, _, _ in races] + ['備考', 'メモ']
HR = 4
header(ws, HR, H, [6, 18, 20, 11, 8, 10, 8, 10, 8] + [18] * len(races) + [40, 30])
for c in range(4, 6):
    ws.cell(HR, c).fill, ws.cell(HR, c).font = MAIN_FILL, MAIN_HEAD
C = {'5000m': ('D', 'E'), '3000mSC': ('F', 'G'), '1500m': ('H', 'I')}
NOTE = 10 + len(races)
r = HR + 1
for a in athletes:
    put(ws, r, 1, NO[id(a)])
    put(ws, r, 2, a['name'], align=LEFT, font=BOLD if a['r5000'] else BASE)
    put(ws, r, 3, a['school'], align=LEFT)
    for ev, (tc, rc) in C.items():
        rk, t = a['ev'].get(ev, (None, None))
        f = BOLD if ev == '5000m' else BASE
        put(ws, r, col_idx(tc), tval(t) if t else None, TFMT, f)
        put(ws, r, col_idx(rc), rk, '0', f)
    for j, (sheet, _, _) in enumerate(races):
        put(ws, r, 10 + j, a.get('races', {}).get(sheet), align=LEFT)
    put(ws, r, NOTE, '／'.join(a['notes']) or None, align=LEFT)
    put(ws, r, NOTE + 1, a.get('memo'), align=LEFT).fill = INPUT_FILL
    mark_row(ws, r, NOTE, a)
    r += 1
LAST = r - 1
rank_colors(ws, f'E{HR + 1}:E{LAST} G{HR + 1}:G{LAST} I{HR + 1}:I{LAST}')
ws.freeze_panes = ws.cell(HR + 1, 4)
ws.auto_filter.ref = f'A{HR}:{get_column_letter(NOTE + 1)}{LAST}'
L = f"'選手一覧'!"
rng = lambda col: f'{L}${col}${HR + 1}:${col}${LAST}'


def other_events(ws, r, c, no_col):
    """選手一覧のNoから5000m・3000mSC・1500mの記録と順位を参照する列を c 列目から書く"""
    for j, ev in enumerate(EVENTS):
        tc, rc = C[ev]
        f = BOLD if ev == '5000m' else BASE
        put(ws, r, c + j * 2, f'=IF(INDEX({rng(tc)},${no_col}{r})="","",INDEX({rng(tc)},${no_col}{r}))', TFMT, f)
        put(ws, r, c + j * 2 + 1, f'=IF(INDEX({rng(rc)},${no_col}{r})="","",INDEX({rng(rc)},${no_col}{r}))', '0', f)


# ---------------- 種目別ランキング（5000mは選手一覧と同じ並びなので作らない） ----------------
for ev in EVENTS[1:]:
    ws = wb.create_sheet(ev)
    # 同一人物が二重掲載されている場合は上位（速い方）の記録だけ残す
    rows = sorted((a for a in athletes if ev in a['ev']), key=lambda a: a['ev'][ev][0])
    others = [o for o in EVENTS if o != ev]
    title(ws, f'{ev} ランキング（陸マガ掲載順）',
          f'記録・順位は元データのまま（同一人物の二重掲載は上位の記録のみ残したため、その順位は欠番）。'
          f'{others[0]}・{others[1]}は「選手一覧」から数式で参照（同一人物と判定できた場合のみ）。')
    H = ['順位', '氏名', '高校', '記録', 'トップ差\n(秒)', 'No', f'{others[0]}\n記録', f'{others[0]}\n順位',
         f'{others[1]}\n記録', f'{others[1]}\n順位']
    header(ws, 4, H, [7, 18, 20, 10, 9, 6, 10, 8, 10, 8])
    r = 5
    for a in rows:
        (rk, t), (name, school) = a['ev'][ev], a['src'][ev]
        put(ws, r, 1, rk, '0')
        put(ws, r, 2, name, align=LEFT)
        put(ws, r, 3, school, align=LEFT)
        put(ws, r, 4, tval(t), TFMT)
        put(ws, r, 5, f'=ROUND((D{r}-$D$5)*86400,2)', '+0.00;-0.00;0')
        put(ws, r, 6, NO[id(a)], '0')
        for j, o in enumerate(others):
            tc, rc = C[o]
            put(ws, r, 7 + j * 2, f'=IF(INDEX({rng(tc)},$F{r})="","",INDEX({rng(tc)},$F{r}))', TFMT)
            put(ws, r, 8 + j * 2, f'=IF(INDEX({rng(rc)},$F{r})="","",INDEX({rng(rc)},$F{r}))', '0')
        mark_row(ws, r, 10, a)
        r += 1
    rank_colors(ws, f'A5:A{r - 1} H5:H{r - 1} J5:J{r - 1}')
    ws.freeze_panes = 'D5'
    ws.auto_filter.ref = f'A4:J{r - 1}'

# ---------------- 駅伝 ----------------
for pos, (sheet, label, runs) in enumerate(races, 1):
    ws = wb.create_sheet(sheet, pos)
    title(ws, f'{label}：2028年入学世代（現高2）の区間成績',
          '全出走者のうち、陸マガランキングに名前と高校が一致する選手のみ掲載（成績表に学年の記載がないため）。'
          'ランキング圏外の2年生は含まれません。区間順位は全学年での順位、2年生内順位はこの表の中での順位。'
          '区間順位の色：濃い緑＝3位以内、緑＝10位以内、薄緑＝30位以内。')
    H = ['区', '距離\n(km)', '区間\n順位', '2年生内\n順位', '区間記録', '1km\nペース', '氏名', 'チーム', '都道府県',
         'チーム\n総合順位', 'No', '5000m\n記録', '5000m\n順位', '3000mSC\n記録', '3000mSC\n順位', '1500m\n記録', '1500m\n順位']
    header(ws, 4, H, [5, 8, 7, 8, 10, 8, 18, 20, 9, 8, 6, 11, 8, 10, 8, 10, 8])
    for c in (12, 13):
        ws.cell(4, c).fill, ws.cell(4, c).font = MAIN_FILL, MAIN_HEAD
    E = 4 + len(runs)
    r = 5
    for x in runs:
        put(ws, r, 1, int(x['section']), '0"区"')
        put(ws, r, 2, float(x['distance_km']), '0.0###')
        put(ws, r, 3, int(x['rank']), '0', BOLD)
        put(ws, r, 4, f'=COUNTIFS($A$5:$A${E},A{r},$E$5:$E${E},"<"&E{r})+1', '0')
        put(ws, r, 5, time_val(x['time']), '[m]:ss')
        put(ws, r, 6, f'=E{r}/B{r}', '[m]:ss')
        put(ws, r, 7, x['name'], align=LEFT)
        put(ws, r, 8, x['team'], align=LEFT)
        put(ws, r, 9, x['pref'] or None)
        put(ws, r, 10, int(x['team_rank']), '0')
        put(ws, r, 11, x['no'], '0')
        other_events(ws, r, 12, 'K')
        r += 1
    rank_colors(ws, f'C5:C{E}', ((3, TOP10), (10, TOP50), (30, TOP100)))
    rank_colors(ws, f'M5:M{E} O5:O{E} Q5:Q{E}')
    prev, on = None, False
    for rr in range(5, E + 1):  # 区が変わるごとに薄い帯色
        if ws.cell(rr, 1).value != prev:
            on, prev = not on, ws.cell(rr, 1).value
        if on:
            for c in (1, 2, 4, 5, 6, 7, 8, 9, 10, 11):
                ws.cell(rr, c).fill = PatternFill('solid', fgColor='F2F2F2')
    for rr, x in enumerate(runs, 5):
        mark_row(ws, rr, 17, athletes[x['no'] - 1])
    if not any(x['pref'] for x in runs):
        ws.column_dimensions['I'].hidden = True  # 都道府県の記載がない大会
    ws.freeze_panes = 'H5'
    ws.auto_filter.ref = f'A4:Q{E}'

# ---------------- 注目選手 ----------------
ws = wb.create_sheet('注目選手', 1)
title(ws, f'注目選手（{len(FOCUS)}人・5000m順）',
      '選手一覧でオレンジにした選手。記録・駅伝・メモは「選手一覧」から数式で参照（メモは選手一覧のメモ列に記入）。'
      '注目選手を増減するときは、どのシートでも行をオレンジ（FFC000）にして返してください。')
H = ['No', '氏名', '高校', '5000m\n記録', '5000m\n順位', '3000mSC\n記録', '3000mSC\n順位', '1500m\n記録', '1500m\n順位']
H += [f'{sheet}\n（区・区間順位）' for sheet, _, _ in races] + ['備考', 'メモ']
header(ws, 4, H, [6, 18, 20, 11, 8, 10, 8, 10, 8] + [18] * len(races) + [40, 30])
for c in range(4, 6):
    ws.cell(4, c).fill, ws.cell(4, c).font = MAIN_FILL, MAIN_HEAD
for r, a in enumerate(FOCUS, 5):
    put(ws, r, 1, NO[id(a)])
    for c in range(2, NOTE + 2):
        col = get_column_letter(c)
        fmt = TFMT if col in 'DFH' else '0' if col in 'EGI' else None
        put(ws, r, c, f'=IF(INDEX({rng(col)},$A{r})="","",INDEX({rng(col)},$A{r}))', fmt,
            BOLD if col in 'BDE' else BASE, LEFT if c in (2, 3) or c >= 10 else CENTER)
FLAST = 4 + len(FOCUS)
rank_colors(ws, f'E5:E{FLAST} G5:G{FLAST} I5:I{FLAST}')
ws.freeze_panes = 'D5'
ws.auto_filter.ref = f'A4:{get_column_letter(NOTE + 1)}{FLAST}'

for ws in wb:
    ws.sheet_view.zoomScale = 90
wb.save(sys.argv[2])
print(len(entries), 'entries ->', len(athletes), 'athletes,', len(FOCUS), 'focus,',
      ', '.join(f'{s} {len(x)}' for s, _, x in races))
